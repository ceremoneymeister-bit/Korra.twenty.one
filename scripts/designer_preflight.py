#!/usr/bin/env python3
"""Offline Designer packaging probe, NOT an evaluation of model output.

Run in a disposable network-none container with current engine mounted read-only
and only --out writable. No credentials, production DATA, gateway or scheduler.
Creates the template through the real HTTP API, then exercises its installed
PowerPoint scripts on a two-slide synthetic fixture. Missing runtime dependencies
are a failed preflight, never a passed Designer acceptance.
"""

from __future__ import annotations

import argparse
import importlib.util
import json
import os
from pathlib import Path
import shlex
import shutil
import subprocess
import sys
import time

ENGINE = Path(__file__).resolve().parents[1]
FIXTURES = Path(__file__).parent / "fixtures" / "designer-acceptance"


def dependencies() -> dict:
    return {
        "python_pptx": importlib.util.find_spec("pptx") is not None,
        "pillow": importlib.util.find_spec("PIL") is not None,
        "soffice": shutil.which("soffice") is not None,
        "pdftoppm": shutil.which("pdftoppm") is not None,
        "pdftotext": shutil.which("pdftotext") is not None,
    }


def command(argv: list[str]) -> str:
    result = subprocess.run(
        argv, capture_output=True, text=True, encoding="utf-8", timeout=180,
        check=True,
    )
    return result.stdout


def powerpoint_probe(profile: Path, output: Path) -> dict:
    """Exercise shipped create/read/render/export, including Cyrillic readback."""
    from PIL import Image
    from pptx import Presentation

    matches = list((profile / "skills").rglob("powerpoint/scripts/pptx_create.py"))
    if len(matches) != 1:
        raise RuntimeError("Expected one installed powerpoint skill")
    scripts = matches[0].parent
    output.mkdir(parents=True, exist_ok=False)
    deck = output / "production-probe.pptx"
    command([sys.executable, str(scripts / "pptx_create.py"),
             str(FIXTURES / "production-probe.json"), str(deck)])
    outline = json.loads(command([sys.executable, str(scripts / "pptx_read.py"),
                                  str(deck), "--outline"]))
    (output / "outline.json").write_text(json.dumps(outline, ensure_ascii=False, indent=2), encoding="utf-8")
    presentation = Presentation(deck)
    texts = [shape.text for slide in presentation.slides for shape in slide.shapes if shape.has_text_frame]
    if len(presentation.slides) != 2 or "Проверка кириллицы" not in texts:
        raise RuntimeError("PPTX slide count or editable Cyrillic text mismatch")
    rendered = json.loads(command([sys.executable, str(scripts / "pptx_render.py"),
                                   str(deck), "--outdir", str(output / "render")]))
    if rendered.get("rendered") is not True or len(rendered.get("files", [])) != 2:
        raise RuntimeError("Render did not produce both slides")
    dimensions = []
    for filename in rendered["files"]:
        with Image.open(filename) as image:
            image.load()
            dimensions.append(list(image.size))
    command(["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(output), str(deck)])
    pdf = deck.with_suffix(".pdf")
    if not pdf.is_file() or not pdf.read_bytes().startswith(b"%PDF-"):
        raise RuntimeError("PDF export missing")
    extracted = command(["pdftotext", str(pdf), "-"])
    if "Проверка кириллицы" not in extracted:
        raise RuntimeError("Cyrillic text lost in PDF export")
    return {"status": "PASS", "slides": 2, "editable_text": True,
            "pdf_cyrillic": True, "render_dimensions": dimensions,
            "visual_review": "NOT_RUN", "designer_generated": False}


def review_board_probe(profile: Path, output: Path) -> dict:
    """Exercise the installed comparison helper without model or user data."""
    from PIL import Image

    matches = list((profile / "skills").rglob("visual-design/scripts/review_board.py"))
    if len(matches) != 1:
        raise RuntimeError("Expected one installed Designer review-board helper")
    output.mkdir(parents=True, exist_ok=False)
    candidate = output / "candidate.png"
    reference = output / "reference.png"
    board = output / "board.png"
    Image.new("RGB", (400, 600), "red").save(candidate)
    Image.new("RGBA", (600, 400), (0, 0, 255, 255)).save(reference)
    before = [path.read_bytes() for path in (candidate, reference)]
    command([
        sys.executable, str(matches[0]),
        "--candidate", str(candidate),
        "--reference", str(reference),
        "--output", str(board),
    ])
    with Image.open(board) as image:
        image.load()
        if image.size != (2400, 2200):
            raise RuntimeError("Designer review board has an unexpected size")
        if image.getpixel((765, 1100)) != (255, 0, 0):
            raise RuntimeError("Designer review board lost the candidate")
        if image.getpixel((1955, 1100)) != (0, 0, 255):
            raise RuntimeError("Designer review board lost the reference")
    if [path.read_bytes() for path in (candidate, reference)] != before:
        raise RuntimeError("Designer review board changed an input image")
    return {"status": "PASS", "size": [2400, 2200], "inputs_preserved": True}


def install_template(home: Path, *, model: str | None = None) -> dict:
    from fastapi.testclient import TestClient
    from agent.prompt_builder import load_soul_md
    from korra_cli.config import load_config
    from korra_cli.tools_config import _get_platform_tools
    from korra_cli.web_server import app
    from korra_constants import set_hermes_home_override, reset_hermes_home_override
    from tools.skills_tool import skills_list, skill_view

    # Use real routes/middleware without dashboard lifespan (background jobs).
    client = TestClient(app)
    try:
        client.headers["Authorization"] = "Bearer designer-offline-preflight"
        catalogue = client.get("/api/agent-templates")
        catalogue.raise_for_status()
        template = next(item for item in catalogue.json()["templates"] if item["id"] == "korra.designer")
        body = {
            "name": "designer", "display_name": "Дизайнер — приёмка",
            "template_id": template["id"], "template_version": template["version"],
            "idempotency_key": "designer-offline-preflight-01",
        }
        if model:
            body.update(provider="openai-codex", model=model)
        response = client.post("/api/profiles", json=body)
        response.raise_for_status()
    finally:
        client.close()
    profile = home / "profiles/designer"
    if Path(response.json()["path"]).resolve() != profile.resolve():
        raise RuntimeError("Profile created outside the test home")
    if "Дизайнер" not in load_soul_md(home_override=profile):
        raise RuntimeError("Designer role not loaded")
    payload = response.json()
    generation = payload.get("generation")
    if generation != {
        "configured": True,
        "available": False,
        "status": "needs_auth",
        "provider": "openai-codex",
        "model": "gpt-image-2.5-sunburst",
        "platforms": ["cli", "api_server"],
        "live_tested": False,
    } or payload.get("generation_checked") is not False:
        raise RuntimeError(
            "Offline Designer must configure GPT Image without claiming a live test"
        )
    scope = set_hermes_home_override(str(profile))
    try:
        config = load_config()
        if config.get("image_gen") != {
            "provider": "openai-codex",
            "model": "gpt-image-2.5-sunburst",
        }:
            raise RuntimeError("Designer image route was not saved in its profile")
        for platform in generation["platforms"]:
            if "image_gen" not in _get_platform_tools(config, platform):
                raise RuntimeError(
                    f"Designer image tool is disabled on the {platform} surface"
                )
        listed = json.loads(skills_list())
        names = {skill["name"] for skill in listed["skills"]}
        if not {"visual-design", "powerpoint"}.issubset(names):
            raise RuntimeError("Required skills not discovered")
        if not json.loads(skill_view("visual-design", preprocess=False))["success"]:
            raise RuntimeError("Designer instructions not readable")
    finally:
        reset_hermes_home_override(scope)
    return {"status": "PASS", "template_id": template["id"], "version": template["version"],
            "model_configured": payload["model_set"], "generation": generation,
            "skill_count": len(names)}


def native_terminal_probe(profile: Path) -> dict:
    """Check the interpreter actually used by the agent, not just this process."""
    from korra_constants import set_hermes_home_override, reset_hermes_home_override
    from tools.terminal_tool import terminal_tool

    scope = set_hermes_home_override(str(profile))
    code = "import json,sys,pptx; from PIL import Image; print(json.dumps(dict(python=sys.executable,prefix=sys.prefix,pptx=pptx.__version__)))"
    try:
        results = []
        for interpreter in ("python", "python3"):
            result = json.loads(terminal_tool(
                command=f"{interpreter} -c {shlex.quote(code)}", workdir=str(profile), timeout=30,
            ))
            if result.get("exit_code") != 0 or result.get("error"):
                raise RuntimeError(f"Native terminal {interpreter} cannot import presentation dependencies: {result.get('output', '')[:1000]}")
            results.append(json.loads(result["output"]))
        return {"status": "PASS", "interpreters": results}
    finally:
        reset_hermes_home_override(scope)


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, required=True, help="New directory inside the disposable mount")
    args = parser.parse_args()
    if os.environ.get("KORRA_DESIGNER_PREFLIGHT") != "isolated" or not Path("/.dockerenv").exists():
        parser.error("Use a disposable container with KORRA_DESIGNER_PREFLIGHT=isolated and --network none")
    output = args.out.resolve()
    output.mkdir(parents=True, exist_ok=False)
    home = output / "home"
    home.mkdir(mode=0o700)
    # Set before importing engine modules; no inherited user profile is read.
    os.environ["HERMES_HOME"] = str(home)
    os.environ["KORRA_HOME"] = str(home)
    os.environ["HERMES_DASHBOARD_SESSION_TOKEN"] = "designer-offline-preflight"
    os.environ["PYTHONDONTWRITEBYTECODE"] = "1"
    os.environ["TERMINAL_ENV"] = "local"
    os.environ["TERMINAL_CWD"] = str(home)
    sys.path.insert(0, str(ENGINE))
    report = {"kind": "offline_designer_preflight", "model_calls": 0,
              "image_calls": 0, "behavioral_acceptance": "NOT_RUN",
              "dependencies": dependencies(),
              "review_board": {"status": "NOT_RUN"},
              "production": {"status": "NOT_RUN"}}
    started = time.monotonic()
    try:
        report["installation"] = install_template(home)
        missing = [name for name, available in report["dependencies"].items() if not available]
        if missing:
            report["status"] = "FAIL"
            report["missing"] = missing
        else:
            report["native_terminal"] = native_terminal_probe(home / "profiles/designer")
            report["review_board"] = review_board_probe(
                home / "profiles/designer", output / "review-board",
            )
            report["production"] = powerpoint_probe(home / "profiles/designer", output / "production")
            report["status"] = "PASS"
    except Exception as exc:
        report["status"] = "FAIL"
        report["error"] = f"{type(exc).__name__}: {exc}"
    report["elapsed_seconds"] = round(time.monotonic() - started, 2)
    (output / "report.json").write_text(json.dumps(report, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return 0 if report["status"] == "PASS" else 1


if __name__ == "__main__":
    raise SystemExit(main())
