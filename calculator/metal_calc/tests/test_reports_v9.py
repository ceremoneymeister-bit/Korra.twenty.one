"""V9 internal/client reports and deterministic book export."""
from __future__ import annotations

import base64
import copy
import hashlib
import json
from io import BytesIO
from pathlib import Path
from typing import Any

import pytest
from openpyxl import load_workbook

from metal_calc.errors import Conflict, InvalidState, PackChanged
from metal_calc.packs2 import ACTIVE_POINTER, PipelinePackStore
from metal_calc.registry import Registry
from metal_calc.securefs import SecureRoot
from metal_calc.service3 import WorkflowService
from metal_calc.util import digest_json, sha256_bytes

from test_workflow_v3 import drive_to_book, workflow_pack


FINAL_ROUTE = [
    {"seq": 1, "process_code": "BLANK.CUTOFF", "execution_mode": "in_house", "note": "распил"},
    {"seq": 2, "process_code": "MACHINING.TURN.CNC", "execution_mode": "in_house", "note": "точение"},
    {"seq": 3, "process_code": "FINISH.FITTER", "execution_mode": "in_house", "note": "слесарка"},
]


@pytest.fixture
def svc(tmp_path: Path) -> WorkflowService:
    rates = tmp_path / "rates3"
    rates.mkdir()
    pack = copy.deepcopy(workflow_pack())
    for norm in pack["norm_params"].values():
        norm["source"]["ref"] = norm["source"]["ref"].replace(" (TEMPLATE)", "")
    pack["overheads"]["source"]["ref"] = pack["overheads"]["source"]["ref"].replace(
        " (TEMPLATE)", ""
    )
    payload = json.dumps(pack, ensure_ascii=False, sort_keys=True).encode("utf-8")
    (rates / "v3-test.json").write_bytes(payload)
    (rates / ACTIVE_POINTER).write_text(
        json.dumps({"revision": "v3-test", "sha256": sha256_bytes(payload)}),
        encoding="utf-8",
    )
    return WorkflowService(
        Registry(tmp_path / "orders3" / "registry.db"),
        PipelinePackStore(SecureRoot(rates, writable=False)),
    )


def drive_to_ready(svc: WorkflowService, *, order_id: str = "ord-v3-1", **kwargs: Any) -> int:
    revision = drive_to_book(svc, order_id, route_steps=FINAL_ROUTE, **kwargs)
    revision = svc.qa_run_mechanical(order_id, revision, actor_role="front")["revision"]
    revision = svc.qa_verdict(
        order_id, revision, "technological", "PASS", [], "методолог", actor_role="qa"
    )["revision"]
    return svc.qa_verdict(
        order_id, revision, "commercial", "PASS", [], "методолог", actor_role="qa"
    )["revision"]


def test_internal_report_has_frozen_input_cost_qa_and_blockers(svc: WorkflowService) -> None:
    drive_to_book(svc, route_steps=FINAL_ROUTE)
    before_revision = svc.registry.get("ord-v3-1")[0]
    report = svc.report_get("ord-v3-1", "internal", actor_role="front")
    assert report["report_version"] == "v9"
    assert report["document_status"] == "PRELIMINARY"
    assert report["input"]["origin"] == "order_registry.source_files"
    assert report["input"]["quantity"] == 10
    assert report["workflow"]["status"] == "BOOK_ASSEMBLED"
    assert report["book"]["cost"]["amount_kind"] == "INTERNAL_COST"
    assert report["book"]["digest"]
    assert report["qa"] == {}
    assert {item["code"] for item in report["blockers"]} == {"qa_incomplete"}
    assert svc.registry.get("ord-v3-1")[0] == before_revision


def test_client_preview_is_allowlisted_and_explicitly_final(svc: WorkflowService) -> None:
    drive_to_ready(svc)
    preview = svc.report_get("ord-v3-1", "client", actor_role="front")
    assert preview["document_status"] == "FINAL"
    assert preview["is_final"] is True
    assert preview["status_text"] == "Окончательная цена"
    assert preview["price"]["amount_kind"] == "CUSTOMER_PRICE"
    serialized = json.dumps(preview, ensure_ascii=False)
    for forbidden in (
        "INTERNAL_COST",
        "pack_fingerprint",
        "rate_source",
        "source_manifest",
        "input_digest",
        "margin",
    ):
        assert forbidden not in serialized


def test_client_preview_never_promotes_book_without_current_qa(svc: WorkflowService) -> None:
    drive_to_book(svc, route_steps=FINAL_ROUTE)
    preview = svc.report_get("ord-v3-1", "client", actor_role="front")
    assert preview["document_status"] == "PRELIMINARY"
    assert preview["is_final"] is False
    assert preview["status_text"] == "Предварительная цена"
    assert preview["blockers"] == [
        {"code": "approval_pending", "message": "Расчёт проходит внутреннюю проверку"}
    ]


def test_expired_price_is_never_final(svc: WorkflowService) -> None:
    drive_to_ready(svc)
    revision, _ = svc.registry.get("ord-v3-1")

    def expire(saved: dict[str, Any]) -> None:
        saved["book"]["price"]["valid_until"] = "2000-01-01"
        saved["book"]["digest"] = digest_json(
            {key: value for key, value in saved["book"].items() if key != "digest"}
        )

    svc.registry.mutate("ord-v3-1", expire, expected_revision=revision)
    preview = svc.report_get("ord-v3-1", "client", actor_role="front")
    assert preview["document_status"] == "PRELIMINARY"
    assert preview["is_final"] is False
    assert {item["code"] for item in preview["blockers"]} == {"price_expired"}


def test_client_preview_rejects_human_forged_mechanical_pass(svc: WorkflowService) -> None:
    drive_to_book(svc, route_steps=FINAL_ROUTE)
    revision, state = svc.registry.get("ord-v3-1")
    calculation_revision = str(state["workflow"]["calculation_revision"])

    def forge_passes(saved: dict[str, Any]) -> None:
        saved["qa"] = {
            calculation_revision: {
                gate: {"verdict": "PASS", "by": "operator"}
                for gate in ("mechanical", "technological", "commercial")
            }
        }
        saved["workflow"]["status"] = "READY_FOR_LD"
        saved["status"] = "READY_FOR_LD"

    svc.registry.mutate("ord-v3-1", forge_passes, expected_revision=revision)
    preview = svc.report_get("ord-v3-1", "client", actor_role="front")
    assert preview["document_status"] == "PRELIMINARY"
    assert preview["is_final"] is False
    assert preview["blockers"] == [
        {"code": "approval_pending", "message": "Расчёт проходит внутреннюю проверку"}
    ]


def test_client_preview_rejects_human_gate_without_qa_provenance(
    svc: WorkflowService,
) -> None:
    drive_to_ready(svc)
    revision, state = svc.registry.get("ord-v3-1")
    calculation_revision = str(state["workflow"]["calculation_revision"])

    def forge_actor(saved: dict[str, Any]) -> None:
        saved["qa"][calculation_revision]["technological"]["actor_role"] = "front"

    svc.registry.mutate("ord-v3-1", forge_actor, expected_revision=revision)
    preview = svc.report_get("ord-v3-1", "client", actor_role="front")
    assert preview["document_status"] == "PRELIMINARY"
    assert preview["is_final"] is False
    assert preview["blockers"] == [
        {"code": "approval_pending", "message": "Расчёт проходит внутреннюю проверку"}
    ]


def test_report_tools_reject_non_front_role(svc: WorkflowService) -> None:
    drive_to_book(svc, route_steps=FINAL_ROUTE)
    with pytest.raises(InvalidState, match="роли front"):
        svc.report_get("ord-v3-1", "internal", actor_role="qa")
    with pytest.raises(InvalidState, match="роли front"):
        svc.render_book_xlsx(
            "ord-v3-1", "a" * 64, actor_role="qa"
        )


def test_book_xlsx_is_byte_stable_exact_and_contains_no_internal_cost(
    svc: WorkflowService,
) -> None:
    drive_to_ready(svc)
    order_revision, state = svc.registry.get("ord-v3-1")
    book_digest = state["book"]["digest"]
    first = svc.render_book_xlsx("ord-v3-1", book_digest, actor_role="front")
    second = svc.render_book_xlsx("ord-v3-1", book_digest, actor_role="front")
    assert first == second
    payload = base64.b64decode(first["content_base64"], validate=True)
    assert first["sha256"] == hashlib.sha256(payload).hexdigest()
    assert first["bytes"] == len(payload)
    assert first["book_digest"] == book_digest
    assert first["document_status"] == "FINAL"
    workbook = load_workbook(BytesIO(payload), data_only=False)
    values = [cell.value for row in workbook.active.iter_rows() for cell in row]
    text = "\n".join(str(value) for value in values if value is not None)
    assert "INTERNAL_COST" not in text
    assert "Себестоимость" not in text
    assert "ОКОНЧАТЕЛЬНАЯ" in text
    assert "calculation" not in state
    saved_revision, saved = svc.registry.get("ord-v3-1")
    assert saved_revision == order_revision
    assert "artifacts" not in saved


def test_book_xlsx_escapes_formula_like_customer_and_kd(svc: WorkflowService) -> None:
    drive_to_book(
        svc,
        route_steps=FINAL_ROUTE,
        customer={"name": '=HYPERLINK("https://example.invalid","x")'},
        kd_revision="=1+1",
    )
    _, state = svc.registry.get("ord-v3-1")
    rendered = svc.render_book_xlsx(
        "ord-v3-1", state["book"]["digest"], actor_role="front"
    )
    workbook = load_workbook(
        BytesIO(base64.b64decode(rendered["content_base64"])), data_only=False
    )
    cells = [cell for row in workbook.active.iter_rows() for cell in row]
    hostile = [cell for cell in cells if isinstance(cell.value, str) and "HYPERLINK" in cell.value]
    assert len(hostile) == 1
    assert hostile[0].value.startswith("'=")
    assert hostile[0].data_type == "s"
    assert all(
        not cell.value.startswith(("=", "+", "-", "@", "\t", "\r"))
        for cell in cells
        if isinstance(cell.value, str)
    )


def test_report_and_xlsx_fail_closed_on_book_digest_mismatch(svc: WorkflowService) -> None:
    drive_to_book(svc, route_steps=FINAL_ROUTE)
    revision, state = svc.registry.get("ord-v3-1")
    exact_digest = state["book"]["digest"]
    with pytest.raises(Conflict, match="digest"):
        svc.render_book_xlsx("ord-v3-1", "f" * 64, actor_role="front")

    def corrupt(saved: dict[str, Any]) -> None:
        saved["book"]["price"]["total_rub"] += 1

    svc.registry.mutate("ord-v3-1", corrupt, expected_revision=revision)
    with pytest.raises(Conflict, match="digest"):
        svc.report_get("ord-v3-1", "internal", actor_role="front")
    with pytest.raises(Conflict, match="digest"):
        svc.render_book_xlsx("ord-v3-1", exact_digest, actor_role="front")


def test_client_and_xlsx_fail_closed_on_stale_pack(svc: WorkflowService, tmp_path: Path) -> None:
    drive_to_book(svc, route_steps=FINAL_ROUTE)
    _, state = svc.registry.get("ord-v3-1")
    changed = workflow_pack("v3-next")
    changed["machines"]["turning"]["rate_rub_per_hour"] = "9999"
    payload = json.dumps(changed, ensure_ascii=False, sort_keys=True).encode("utf-8")
    rates = tmp_path / "rates3"
    (rates / "v3-next.json").write_bytes(payload)
    (rates / ACTIVE_POINTER).write_text(
        json.dumps({"revision": "v3-next", "sha256": sha256_bytes(payload)}),
        encoding="utf-8",
    )
    internal = svc.report_get("ord-v3-1", "internal", actor_role="front")
    assert "pack_changed" in {item["code"] for item in internal["blockers"]}
    with pytest.raises(PackChanged):
        svc.report_get("ord-v3-1", "client", actor_role="front")
    with pytest.raises(PackChanged):
        svc.render_book_xlsx(
            "ord-v3-1", state["book"]["digest"], actor_role="front"
        )


def test_stale_book_calculation_revision_is_rejected(svc: WorkflowService) -> None:
    drive_to_book(svc, route_steps=FINAL_ROUTE)
    revision, state = svc.registry.get("ord-v3-1")
    digest = state["book"]["digest"]

    def make_stale(saved: dict[str, Any]) -> None:
        saved["workflow"]["calculation_revision"] += 1

    svc.registry.mutate("ord-v3-1", make_stale, expected_revision=revision)
    with pytest.raises(Conflict, match="текущей ревизии расчёта"):
        svc.render_book_xlsx("ord-v3-1", digest, actor_role="front")


def test_frozen_v9_input_rejects_a_late_attachment(service, tmp_path: Path) -> None:
    rates = tmp_path / "rates-late"
    rates.mkdir()
    pack = workflow_pack()
    payload = json.dumps(pack, ensure_ascii=False, sort_keys=True).encode("utf-8")
    (rates / "v3-test.json").write_bytes(payload)
    (rates / ACTIVE_POINTER).write_text(
        json.dumps({"revision": "v3-test", "sha256": sha256_bytes(payload)}),
        encoding="utf-8",
    )
    workflow = WorkflowService(
        service.registry,
        PipelinePackStore(SecureRoot(rates, writable=False)),
    )
    order_id = "late-file"
    service.order_upsert(order_id, 0, {"customer": {"name": "Клиент"}})
    first_cache = "doc_aaaaaaaaaaaa_first.pdf"
    second_cache = "doc_bbbbbbbbbbbb_second.pdf"
    (service.settings.cache_root / first_cache).write_bytes(b"%PDF-1.4\nfirst")
    (service.settings.cache_root / second_cache).write_bytes(b"%PDF-1.4\nsecond")
    first = service.ingest_attachment(order_id, first_cache)
    revision = service.order_get(order_id)["revision"]
    workflow.input_freeze(
        order_id,
        revision,
        1,
        "R1",
        [
            {
                "source_file_id": first["source_file_id"],
                "name": first["name"],
                "sha256": first["sha256"],
            }
        ],
        "Комплект загружен",
        actor_role="front",
    )

    duplicate = service.ingest_attachment(order_id, first_cache)
    assert duplicate["idempotent"] is True
    with pytest.raises(InvalidState, match="уже зафиксирован"):
        service.ingest_attachment(order_id, second_cache)
