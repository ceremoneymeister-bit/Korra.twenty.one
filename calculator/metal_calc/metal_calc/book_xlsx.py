"""Deterministic client XLSX renderer for a V9 workflow book."""
from __future__ import annotations

import re
from datetime import datetime
from io import BytesIO
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

_FIXED_TIME = datetime(1980, 1, 1)


def render_client_book(preview: dict[str, Any]) -> tuple[bytes, int]:
    """Render only the public preview allowlist; no internal book is accepted."""
    if preview.get("mode") != "client" or preview.get("report_version") != "v9":
        raise ValueError("V9 client preview is required")
    price = preview["price"]
    status = preview["document_status"]
    rows: list[list[Any]] = [
        ["Коммерческое предложение"],
        ["Заказ", preview["order_id"]],
        ["Клиент", preview["customer"]["name"]],
        ["Редакция КД", preview["kd_revision"]],
        ["Количество, шт.", preview["quantity"]],
        ["Статус цены", "ОКОНЧАТЕЛЬНАЯ" if status == "FINAL" else "ПРЕДВАРИТЕЛЬНАЯ"],
        ["Пояснение", preview["status_text"]],
        [],
        ["Цена без НДС, руб.", price["net_total_rub"]],
        ["Ставка НДС, %", price["vat_rate_pct"]],
        ["Сумма НДС, руб.", price["vat_amount_rub"]],
        ["Итого, руб.", price["total_rub"]],
        ["Действительно до", price["valid_until"]],
    ]
    blockers = preview.get("blockers") or []
    if blockers:
        rows.extend([[], ["Что ещё уточняется"]])
        rows.extend([[item["message"]] for item in blockers])

    workbook = Workbook()
    workbook.properties.created = _FIXED_TIME
    workbook.properties.modified = _FIXED_TIME
    workbook.properties.creator = None
    workbook.properties.lastModifiedBy = None
    sheet = workbook.active
    sheet.title = "КП"
    for row in rows:
        sheet.append([_safe_cell(value) if isinstance(value, str) else value for value in row])

    sheet["A1"].font = Font(size=16, bold=True)
    for cell in sheet[6]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAD3" if status == "FINAL" else "FFF2CC")
        cell.alignment = Alignment(horizontal="center")
    for row_number in (9, 10, 11, 12, 13):
        sheet.cell(row_number, 1).font = Font(bold=row_number == 12)
    sheet.column_dimensions["A"].width = 34
    sheet.column_dimensions["B"].width = 24
    sheet.freeze_panes = "A9"

    output = BytesIO()
    workbook.save(output)
    return _canonicalize_xlsx(output.getvalue()), len(rows)


def _safe_cell(value: str) -> str:
    value = str(value)
    if value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def _canonicalize_xlsx(payload: bytes) -> bytes:
    """Normalize timestamps/member order so identical input yields identical bytes."""
    timestamp = b"1980-01-01T00:00:00Z"
    source = BytesIO(payload)
    target = BytesIO()
    with ZipFile(source, "r") as archive, ZipFile(
        target, "w", compression=ZIP_DEFLATED, compresslevel=9
    ) as output:
        for name in sorted(archive.namelist()):
            original = archive.getinfo(name)
            member = archive.read(name)
            if name == "docProps/core.xml":
                for field in (b"created", b"modified"):
                    pattern = (
                        rb"(<dcterms:" + field + rb"\b[^>]*>)[^<]*(</dcterms:"
                        + field
                        + rb">)"
                    )
                    member, replacements = re.subn(
                        pattern, rb"\g<1>" + timestamp + rb"\g<2>", member
                    )
                    if replacements != 1:
                        raise ValueError(
                            f"XLSX core property {field.decode('ascii')} is missing or duplicated"
                        )
            info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = original.external_attr
            info.create_system = original.create_system
            output.writestr(info, member, compress_type=ZIP_DEFLATED, compresslevel=9)
    return target.getvalue()


__all__ = ["render_client_book"]
