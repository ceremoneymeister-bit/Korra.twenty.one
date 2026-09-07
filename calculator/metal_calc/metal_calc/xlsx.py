from __future__ import annotations

from datetime import UTC, datetime
from io import BytesIO
import re
from typing import Any
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill


def render_quote(order: dict[str, Any]) -> tuple[bytes, int]:
    workbook = Workbook()
    calculated_at = datetime.fromisoformat(
        order["calculation"]["calculated_at"].replace("Z", "+00:00")
    ).astimezone(UTC).replace(tzinfo=None)
    workbook.properties.created = calculated_at
    workbook.properties.modified = calculated_at
    workbook.properties.creator = "metal_calc_mcp"
    workbook.properties.lastModifiedBy = "metal_calc_mcp"
    sheet = workbook.active
    sheet.title = "КП"
    rows: list[list[Any]] = [
        ["Коммерческое предложение"],
        ["Заказ", order["order_id"]],
        ["Клиент", _safe_cell(order["customer"]["name"])],
        ["Материал", _safe_cell(order["material"]["grade"])],
        ["Толщина, мм", order["material"]["thickness_mm"]],
        [],
        ["Операция", "Количество", "Ед.", "Ставка, руб.", "Сумма, руб."],
    ]
    for line in order["cost"]["lines"]:
        rows.append(
            [
                _safe_cell(line["code"]),
                line["quantity"],
                line["unit"],
                line["rate_rub"],
                line["subtotal_rub"],
            ]
        )
    rows.extend(
        [
            [],
            ["Себестоимость", None, None, None, order["cost"]["total_rub"]],
            ["Цена без НДС", None, None, None, order["price"]["net_total_rub"]],
            ["Ставка НДС, %", order["price"].get("vat_rate_pct", 0)],
            ["Сумма НДС", None, None, None, order["price"]["vat_amount_rub"]],
            ["Итого с НДС", None, None, None, order["price"]["total_rub"]],
            ["Действительно до", order["price"]["valid_until"]],
        ]
    )
    for row in rows:
        sheet.append(row)
    sheet["A1"].font = Font(size=16, bold=True)
    for cell in sheet[7]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill("solid", fgColor="D9EAD3")
        cell.alignment = Alignment(horizontal="center")
    sheet.column_dimensions["A"].width = 34
    for column in ("B", "C", "D", "E"):
        sheet.column_dimensions[column].width = 18
    sheet.freeze_panes = "A8"
    output = BytesIO()
    workbook.save(output)
    return _canonicalize_xlsx(output.getvalue(), calculated_at), len(rows)


def _safe_cell(value: str) -> str:
    value = str(value)
    if value.startswith(("=", "+", "-", "@", "\t", "\r")):
        return "'" + value
    return value


def _canonicalize_xlsx(payload: bytes, calculated_at: datetime) -> bytes:
    """Make XLSX byte-stable so crash recovery can adopt a published artifact."""
    timestamp = calculated_at.replace(microsecond=0).isoformat().encode("ascii") + b"Z"
    source = BytesIO(payload)
    target = BytesIO()
    with ZipFile(source, "r") as archive, ZipFile(
        target, "w", compression=ZIP_DEFLATED, compresslevel=9
    ) as output:
        for name in sorted(archive.namelist()):
            original = archive.getinfo(name)
            member = archive.read(name)
            if name == "docProps/core.xml":
                # openpyxl overwrites ``modified`` with wall-clock time inside
                # save_workbook(). Normalize both fields after save so a retry
                # renders the exact artifact that may already be published.
                for field in (b"created", b"modified"):
                    pattern = (
                        rb"(<dcterms:" + field + rb"\b[^>]*>)[^<]*(</dcterms:"
                        + field
                        + rb">)"
                    )
                    member, replacements = re.subn(
                        pattern, rb"\g<1>" + timestamp + rb"\g<2>", member
                    )
                    if replacements != 1:
                        raise ValueError(
                            f"XLSX core property {field.decode('ascii')} is missing or duplicated"
                        )
            info = ZipInfo(name, date_time=(1980, 1, 1, 0, 0, 0))
            info.compress_type = ZIP_DEFLATED
            info.external_attr = original.external_attr
            info.create_system = original.create_system
            output.writestr(info, member, compress_type=ZIP_DEFLATED, compresslevel=9)
    return target.getvalue()
